#!/usr/bin/env python3
"""Generates docs/ARCHITECTURE.pdf (max 2 pages) and docs/PRESENTATION.pptx (exactly 5 slides)."""

import os
import sys

def generate_pdf():
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether

    pdf_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "ARCHITECTURE.pdf")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=6
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=8,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#334155'),
        spaceAfter=4
    )
    code_style = ParagraphStyle(
        'DocCode',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=4
    )

    story = []

    # PAGE 1: System Overview & Architecture Flow
    story.append(Paragraph("AegisGuard-ULPF — Sovereign Telemetry Architecture", title_style))
    story.append(Paragraph("<b>National Technical Research Organisation (NTRO) · SIH 2026 PS 26156</b>", body_style))
    story.append(Spacer(1, 4))

    story.append(Paragraph("1. Executive Technical Summary", h2_style))
    story.append(Paragraph(
        "AegisGuard-ULPF is a sovereign, air-gappable pre-SIEM log pre-processing pipeline. It ingests heterogeneous perimeter "
        "streams (Syslog UDP/TCP/TLS, CEF, LEEF, Key-Value, JSON, TSV, CSV), guarantees lossless Base64 wire preservation alongside "
        "pre-transformation SHA-256 digests, normalizes events into pinned OCSF v1.1.0 schemas, enforces character-exact byte-span lineage, "
        "and establishes mathematical non-repudiation via RFC 6962 Merkle trees with Ed25519-signed checkpoints.",
        body_style
    ))

    story.append(Paragraph("2. End-to-End Ingestion & Processing Flowchart", h2_style))
    flowchart_text = """
    [Perimeter Appliances: Cisco, Fortinet, Checkpoint, Juniper, Linux, Suricata, Zeek, pfSense]
                                        │ (Syslog RFC 3164/5424/5425, UDP/TCP/TLS, JSON)
                                        ▼
    ┌────────────────────────────────────────────────────────────────────────────────────────┐
    │ 1. INGESTION GATEWAY : Non-blocking queues, backpressure watermarking, fail-closed     │
    └───────────────────────────────────┬────────────────────────────────────────────────────┘
                                        │
                                        ▼
    ┌────────────────────────────────────────────────────────────────────────────────────────┐
    │ 2. LOSSLESS WIRE VAULT : Raw Base64 + SHA-256 wire digest (storage/lossless_archive)   │
    └───────────────────────────────────┬────────────────────────────────────────────────────┘
                                        │
                                        ▼
    ┌────────────────────────────────────────────────────────────────────────────────────────┐
    │ 3. DECLARATIVE PARSING & OCSF MAP : Hot-reloadable Source Packs + Unmapped Bag         │
    │    4. BYTE-SPAN LINEAGE ENGINE    : Character coordinates [start:end] matching wire    │
    │    5. SOVEREIGN PII REDACTION     : Verhoeff D5 Aadhaar, PAN, Luhn IMEI, Email         │
    └───────────────────────────────────┬────────────────────────────────────────────────────┘
                                        │
                                        ▼
    ┌────────────────────────────────────────────────────────────────────────────────────────┐
    │ 6. CRYPTOGRAPHIC INTEGRITY LEDGER : RFC 6962 Merkle Tree (0x00 leaf / 0x01 interior)   │
    │    7. CHECKPOINT NOTARY           : Ed25519 Enclave Signature over Merkle Roots        │
    └───────────────────────────────────┬────────────────────────────────────────────────────┘
                                        │
                                        ▼
    ┌────────────────────────────────────────────────────────────────────────────────────────┐
    │ 8. ENTERPRISE INTEGRATION SINKS   : Parquet (15-col contract), NDJSON, Splunk, Elastic │
    └────────────────────────────────────────────────────────────────────────────────────────┘
    """
    story.append(Paragraph(flowchart_text.replace(" ", "&nbsp;").replace("\n", "<br/>"), code_style))

    story.append(Paragraph("3. Core Subsystems & Cryptographic Contracts", h2_style))
    summary_data = [
        ["Subsystem", "Wire Protocol / Spec", "Cryptographic Guarantee", "Operational Boundary"],
        ["Lossless Archive", "JSONL / Base64", "SHA-256 pre-mutation digest", "100% byte retention; 0% dropped"],
        ["Merkle Engine", "RFC 6962 CT-style", "0x00 leaf / 0x01 interior hash", "Logarithmic O(log N) inclusion proofs"],
        ["Checkpoint Notary", "Ed25519 Elliptic Curve", "Public key verifiable signature", "Catches root re-sealing attacks"],
        ["Lineage Engine", "Character offset slices", "[start:end] exact character match", "Math proof against token corruption"]
    ]
    t = Table(summary_data, colWidths=[100, 120, 160, 160])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
        ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor('#0F172A')),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 7.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    story.append(t)

    # PAGE 2: Requirements, Empirical Metrics & Known Limitations
    story.append(PageBreak())

    story.append(Paragraph("4. Official SIH 26156 Requirement Traceability (Items a - k)", h2_style))
    req_data = [
        ["Item", "Requirement Statement", "AegisGuard-ULPF Verification Mechanism", "Status"],
        ["(a)", "Lossless raw event preservation", "Base64 wire payload + SHA-256 digest in lossless_archive.jsonl", "PASS"],
        ["(b)", "Extract source attributes", "Multi-vendor declarative packs (Syslog, KV, CEF, LEEF, JSON, TSV)", "PASS"],
        ["(c)", "Normalize common taxonomy", "Strictly pinned OCSF v1.1.0 schema (Classes 4001, 3002, 2001, 1001)", "PASS"],
        ["(d)", "Traceability & byte lineage", "FieldLineageSpan coordinates [start, end] matching raw string slices", "PASS"],
        ["(e)", "Plug-and-play onboarding", "Adaptive Source Intelligence offline clustering & candidate generator", "PASS"],
        ["(f)", "Unified enterprise visibility", "Live React Cyber Console, category doughnut, threat attribution radar", "PASS"],
        ["(g)", "SIEM & Data Lake integration", "Parquet 15-col contract, NDJSON, Splunk HEC, Elastic bulk forwarders", "PASS"],
        ["(h)", "AI/ML-ready analytics", "Typed OCSF schema enriched with local offline APT threat intel IOCs", "PASS"],
        ["(i)", "Reduced parser development", "Deterministic offline drafter (<15 ms) with human approval gate", "PASS"],
        ["(j)", "Air-gapped deployable", "Kernel socket interceptor (EPERM), zero cloud calls, zero egress", "PASS"],
        ["(k)", "Container packaged", "Multi-stage Dockerfile, docker-compose, offline tar bundle script", "PASS"]
    ]
    t2 = Table(req_data, colWidths=[35, 150, 310, 45])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
        ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor('#0F172A')),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 7),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2.5),
        ('TOPPADDING', (0,0), (-1,-1), 2.5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    story.append(t2)
    story.append(Spacer(1, 6))

    story.append(Paragraph("5. Empirical Measured Performance & Operational Boundaries", h2_style))
    story.append(Paragraph(
        "<b>Hardware Testbed:</b> AMD PRO A4-3350B (4 Cores @ 2.0 GHz), 3.3 GB RAM, Linux x86_64, Python 3.14.6.<br/>"
        "• <b>Measured Throughput:</b> Pure in-memory parsing > 100,000 EPS. Streaming pipeline (routing + lineage + Merkle): 11,030 EPS. "
        "Full SQLite WAL persistence: 4,493 EPS (single-worker) / 2,809 EPS (2-worker disk constrained).<br/>"
        "• <b>Scale Arithmetic:</b> 1 Billion events/day requires 11,574 EPS sustained. Single Python worker reaches ~4,500 EPS with full WAL writes; "
        "horizontal scaling across stateless workers (Kafka/Redpanda partitioned design) is required for 1B/day sustained.<br/>"
        "• <b>Security Hardening:</b> Bearer token authentication (0600 file perms), /health exemption, RFC 5425 TLS syslog listener, "
        "and ReDoS catastrophic backtracking static rejection.",
        body_style
    ))

    doc.build(story)
    print(f"[✓] Generated 2-page architecture specification: {pdf_path}")


def generate_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625) # 16:9 widescreen

    DARK_BG = RGBColor(15, 23, 42)    # Slate 900
    TEXT_LIGHT = RGBColor(241, 245, 249)
    TEXT_MUTED = RGBColor(148, 163, 184)
    ACCENT_BLUE = RGBColor(56, 189, 248)
    CARD_BG = RGBColor(30, 41, 59)

    slides_content = [
        # SLIDE 1: Problem
        (
            "NTRO SIH 26156: Universal Log Pre-processing",
            "The Problem: Perimeter Telemetry Blindspots & Forensics Breakdown",
            [
                "Disparate Appliance Schemas: Firewalls, WAFs, and proxies emit incompatible formats (CEF, Syslog, KV, JSON, TSV).",
                "Evidentiary Corruption: Legacy pre-processors mutate strings in-place without raw SHA-256 wire digests or byte lineage.",
                "Air-Gap Violations: SaaS/cloud parsers leak sensitive defense network topology and violate sovereignty mandates.",
                "Scale Challenge: Defense SOCs handle tens of millions to 1B events/day (11,574 EPS sustained) demanding verified architectures."
            ]
        ),
        # SLIDE 2: Solution & Architecture
        (
            "AegisGuard-ULPF Architecture",
            "Sovereign, Lossless, Air-Gapped Telemetry Engine",
            [
                "Lossless Wire Vault: Every raw event Base64 encoded + SHA-256 hashed before transformation (0% dropped).",
                "Pinned Taxonomy: Normalizes multi-vendor streams into strict OCSF v1.1.0 records with unmapped attribute preservation.",
                "RFC 6962 Merkle Tree: Domain-separated cryptographic ledger with Ed25519-signed enclave checkpoints.",
                "Multi-Sink Lakehouse: Native Apache Parquet export with fixed 15-column schema, NDJSON, Splunk HEC, and Elastic bulk."
            ]
        ),
        # SLIDE 3: Key Differentiators & Measured Evidence
        (
            "Key Technical Differentiators",
            "Empirical Evidence vs Industry Prior Art",
            [
                "Bit-Exact Byte Lineage: FieldLineageSpan maps normalized fields to exact [start:end] character slices in raw wire payload.",
                "Empirically Measured Benchmarks: 11,030 EPS streaming; 4,493 EPS full SQLite WAL disk persistence on commodity 4-core hardware.",
                "Real-World Corpora Tested: 100% of 10,000 real Loghub records (OpenSSH, Linux, Apache, Proxifier, HDFS) emitted as OCSF.",
                "Sovereign Indian Compliance: Verhoeff D5 Aadhaar checksum scrubber, Income Tax PAN validator, Luhn IMEI, and CERT-In 6-hour export."
            ]
        ),
        # SLIDE 4: Live Demonstration Flow
        (
            "Two-Minute Evaluation & Demo Flow",
            "Adversarial Verification for Technical Evaluators",
            [
                "Single-Command Launch: ./start_demo.sh brings up offline FastAPI backend, Vite console, and live multi-source syslog.",
                "Forensic Log Drawer: Side-by-side Raw Wire vs OCSF comparison, Base64 payload, and coordinate byte-span pointers.",
                "Tamper Demonstration: Mutate 1 stored byte -> verify_bundle.py rejects; re-seal attack caught by external Ed25519 checkpoint.",
                "Air-Gap Proof: scripts/airgap_proof.sh enforces fail-closed EPERM socket blocking with zero cloud egress."
            ]
        ),
        # SLIDE 5: Scale, Capacity & Known Limitations
        (
            "Scale, Capacity & Engineering Boundaries",
            "Honest Disclosure of Operational Boundaries",
            [
                "The 1B/Day Reality: 1,000,000,000 / 86,400s = 11,574 EPS. Single Python worker reaches ~4,500 EPS with full WAL writes.",
                "Kafka Partitioned Design: Horizontal multi-worker scaling behind partitioned message queues designed for 50k+ EPS.",
                "Software-Enforced Immutability: Software append-only WAL is tamper-evident; physical write protection requires optical media / hardware WORM.",
                "Human Gate on Unknown Onboarding: Deterministic drafter profiles unknown syntax, but requires human operator review before promotion."
            ]
        )
    ]

    for title, subtitle, bullets in slides_content:
        slide_layout = prs.slide_layouts[6] # Blank
        slide = prs.slides.add_slide(slide_layout)

        # Background
        bg_shape = slide.shapes.add_shape(1, 0, 0, Inches(10), Inches(5.625)) # 1 = msoShapeRectangle
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = DARK_BG
        bg_shape.line.color.rgb = DARK_BG

        # Title box
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.4), Inches(1.2))
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Helvetica"
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_BLUE

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = "Helvetica"
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = TEXT_LIGHT

        # Body card
        card = slide.shapes.add_shape(1, Inches(0.8), Inches(1.8), Inches(8.4), Inches(3.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(51, 65, 85)

        card_tf = card.text_frame
        card_tf.word_wrap = True

        for idx, bullet in enumerate(bullets):
            p = card_tf.paragraphs[0] if idx == 0 else card_tf.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.name = "Helvetica"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_LIGHT
            p.space_after = Pt(10)

    pptx_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "PRESENTATION.pptx")
    prs.save(pptx_path)
    print(f"[✓] Generated 5-slide presentation: {pptx_path}")


if __name__ == "__main__":
    generate_pdf()
    generate_pptx()
